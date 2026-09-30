"""Task 9 — standalone-инжест STAINWOOD-документов.

Покрывает:
  1.  extract_txt
  2.  extract_docx
  3.  extract_xlsx (2 sheets → 2 StandaloneDocument, sheet=ws.title)
  4.  extract_doc (OLE2)
  5.  extract_pdf
  6.  extract_file unsupported → ExtractError
  7.  scan_standalone_files — excludes junk, sorts
  8.  ingest_standalone_dir — stats
  9.  standalone_to_chunks — id/stable, doc_type, no source_file → ValueError
 10.  build_index (chunks + stats)
 11.  Retriever returns standalone when product/technology match (FakeModel)
 12.  ContextBuilder renders Title/Source (standalone)
 13.  create_rag_pipeline (standalone_root + standalone_documents)
 14.  End-to-end: FakeLLM + ContextBuilder + standalone chunk in context
"""
from __future__ import annotations

import re
import textwrap
from pathlib import Path

import pytest

from importers.standalone_documents import (
    ExtractError,
    extract_doc,
    extract_docx,
    extract_file,
    extract_pdf,
    extract_txt,
    extract_xlsx,
    ingest_standalone_dir,
    scan_standalone_files,
)
from paint_rag.knowledge.product_store import ProductStore
from paint_rag.models.standalone import StandaloneDocument
from paint_rag.rag.answer_generator import AnswerGenerator
from paint_rag.rag.context_builder import ContextBuilder
from paint_rag.rag.embedding_adapter import ProviderAsModel
from paint_rag.rag.embedding_provider import FakeEmbeddingProvider
from paint_rag.rag.indexing import build_index
from paint_rag.rag.llm import FakeLLM
from paint_rag.rag.pipeline import (
    create_rag_pipeline,
    load_standalone_documents,
)
from paint_rag.rag.retriever import Retriever
from paint_rag.rag.standalone_documents import standalone_to_chunks
from paint_rag.rag.vector_store import VectorStore


DATA = Path("data/knowledge/products.json")


# ------------------------------------------------------------------
# Fixtures (tiny, self-contained)
# ------------------------------------------------------------------


def _write(path: Path, content: str | bytes) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content if isinstance(content, bytes) else content.encode("utf-8"))
    return path


def _make_docx(path: Path) -> Path:
    """Минимальный docx: 2 paragraphs + 1 2x2 table (no sdt)."""
    import zipfile

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    doc_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W}">
<w:body>
<w:p><w:r><w:t>STAINWOOD Method: primer on MDF</w:t></w:r></w:p>
<w:tbl>
  <w:tr><w:tc><w:p><w:r><w:t>base</w:t></w:r></w:p></w:tc>
        <w:tc><w:p><w:r><w:t>top</w:t></w:r></w:p></w:tc></w:tr>
  <w:tr><w:tc><w:p><w:r><w:t>NC</w:t></w:r></w:p></w:tc>
        <w:tc><w:p><w:r><w:t>PE</w:t></w:r></w:p></w:tc></w:tr>
</w:tbl>
<w:p><w:r><w:t>Смешивание: 10:1</w:t></w:r></w:p>
</w:body>
</w:document>
"""

    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/>'
        "</Relationships>"
    )

    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc_xml)
    return path


def _make_xlsx(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "SheetA"
    a["A1"] = "alpha"
    a["B1"] = "1"
    b = wb.create_sheet("SheetB")
    b["A1"] = "beta"
    b["B1"] = "2"
    wb.save(str(path))
    wb.close()
    return path


@pytest.fixture
def tree(tmp_path: Path):
    _write(tmp_path / "note.txt", "Hello\nWorld\nPrimer: Rupa PA777-9016\n")
    _make_docx(tmp_path / "primer.docx")
    _make_xlsx(tmp_path / "ratio.xlsx")
    _write(tmp_path / "junk.bin", "nope")
    _write(tmp_path / "~$temp.docx", b"junk")
    return tmp_path


# ------------------------------------------------------------------
# 1-5. extractors
# ------------------------------------------------------------------


def test_extract_txt_round_trip():
    p = Path("/tmp/sd_test.txt")
    p.write_text("Применение: на МДФ\n", encoding="utf-8")
    docs = extract_txt(p)
    assert len(docs) == 1
    d = docs[0]
    assert d.kind == "txt"
    assert d.source_file
    assert "МДФ" in d.text
    assert d.title


def test_extract_docx_round_trip():
    p = Path("/tmp/sd_test.docx")
    _make_docx(p)
    docs = extract_docx(p)
    assert len(docs) == 1
    d = docs[0]
    assert d.kind == "docx"
    # title = first meaningful line
    assert "STAINWOOD" in (d.title or "")
    # table content present
    assert "NC | PE" in d.text
    assert "alpha" not in d.text  # xlsx content must not leak


def test_extract_xlsx_two_sheets():
    p = Path("/tmp/sd_test.xlsx")
    _make_xlsx(p)
    docs = extract_xlsx(p)
    assert len(docs) == 2
    by_sheet = {d.sheet: d for d in docs}
    assert set(by_sheet) == {"SheetA", "SheetB"}
    assert "alpha" in by_sheet["SheetA"].text
    assert "beta" in by_sheet["SheetB"].text


def test_extract_file_dispatch_by_extension(tree):
    docs = extract_file(tree / "note.txt")
    assert docs and docs[0].kind == "txt"


def test_extract_file_unsupported_raises(tree):
    with pytest.raises(ExtractError):
        extract_file(tree / "junk.bin")


def test_extract_doc_missing_or_bogus(tmp_path: Path, monkeypatch):
    # No .doc fixture is available here: a non-OLE file must still raise
    # a clear ExtractError (not a crash / traceback).
    p = tmp_path / "bogus.doc"
    p.write_bytes(b"not really a Word file")
    with pytest.raises(ExtractError):
        extract_doc(p)


def test_extract_pdf_missing_pypdf_or_bogus(tmp_path: Path, monkeypatch):
    # pypdf may or may not be installed; if it is, a bogus file must
    # still produce ExtractError (not a crash). If pypdf is not importable,
    # the call still must raise (ImportError → treated as "unsupported").
    p = tmp_path / "bogus.pdf"
    p.write_bytes(b"%PDF-1.4 not-really-valid")
    try:
        extract_pdf(p)
        raised = None
    except Exception as exc:  # noqa: BLE001
        raised = exc
    # We do not want a silent success; allow ExtractError / other exceptions.
    # This test asserts the *interface* — the function must raise.
    assert raised is not None, "extract_pdf must raise on unparseable input"


# ------------------------------------------------------------------
# 7. scan_standalone_files
# ------------------------------------------------------------------


def test_scan_standalone_files_excludes_junk_and_sorts(tree):
    files, skipped = scan_standalone_files(tree)
    names = [f.name for f in files]
    assert "junk.bin" not in names
    assert "~$temp.docx" not in names
    assert "note.txt" in names
    assert "primer.docx" in names
    # skipped reasons are recorded
    reasons = {(Path(p).name, r) for p, r in skipped}
    assert any(r.startswith("unsupported ext") for _, r in reasons)


def test_scan_standalone_files_missing_root_raises():
    with pytest.raises(FileNotFoundError):
        scan_standalone_files("/nonexistent-standalone-root-xyzzy")


# ------------------------------------------------------------------
# 8. ingest_standalone_dir
# ------------------------------------------------------------------


def test_ingest_standalone_dir_stats(tree):
    stats = ingest_standalone_dir(tree)
    assert stats["files_extracted"] >= 3
    assert stats["files_error"] == 0
    assert len(stats["documents"]) >= 3
    # by_extension covers txt, docx, xlsx
    assert ".txt" in stats["by_extension"]
    assert ".docx" in stats["by_extension"]
    assert ".xlsx" in stats["by_extension"]
    # every document has a source_file (chunk с источником)
    for d in stats["documents"]:
        assert d.source_file


# ------------------------------------------------------------------
# 9. standalone_to_chunks
# ------------------------------------------------------------------


def test_standalone_to_chunks_basic():
    d = StandaloneDocument(source_file="docs/primer.txt", kind="txt", text="x" * 400)
    chunks = standalone_to_chunks([d], chunk_size=200, overlap=20)
    assert len(chunks) >= 2
    for i, c in enumerate(chunks):
        assert c.doc_type == "standalone"
        assert c.product is None
        assert c.article is None
        assert c.id.startswith("standalone:docs/primer.txt")
        assert c.source["file"] == "docs/primer.txt"
        assert c.source["doc_type"] == "standalone"


def test_standalone_to_chunks_no_source_raises():
    d = StandaloneDocument(source_file="", kind="txt", text="hello")
    with pytest.raises(ValueError, match="source_file"):
        standalone_to_chunks([d], chunk_size=100, overlap=0)


def test_standalone_to_chunks_page_sheet_inclusion():
    d1 = StandaloneDocument(
        source_file="d.pdf", kind="pdf", page=2, text="aaa",
    )
    d2 = StandaloneDocument(
        source_file="m.xlsx", kind="xlsx", sheet="Металл", text="bbb",
    )
    c1, c2 = standalone_to_chunks([d1, d2], chunk_size=1000, overlap=0)
    assert c1.source["page"] == 2
    assert c2.source["sheet"] == "Металл"
    # stable IDs differ
    assert c1.id != c2.id


# ------------------------------------------------------------------
# 10. build_index integrates standalone into one index
# ------------------------------------------------------------------


def test_build_index_includes_standalone():
    store = ProductStore.from_json(DATA)
    provider = FakeEmbeddingProvider(16)
    model = ProviderAsModel(provider)

    standalone = StandaloneDocument(
        source_file="docs/stand.txt", kind="txt", text="Шлифование МДФ до P180",
    )

    vs, stats = build_index(
        store, model, standalone_documents=[standalone],
    )
    assert stats.standalone_documents == 1
    assert stats.standalone_chunks >= 1
    assert stats.total_chunks > stats.chunks - stats.standalone_chunks
    # every chunk has a vector
    assert len(vs.all_chunks()) == stats.total_chunks


class _KeywordModel:
    """Детерминированный keyword-based embed (на подобии FakeModel).

    Standalone-текст с одним из known tokens всегда выигрывает по
    cosine у product-chunks (у которых все биты = 0).
    """

    _TOKENS = ("mdftoken", "uniquestandalone999", "unique999")

    def _vec(self, text: str) -> list[float]:
        low = text.lower()
        return [1.0 if t in low else 0.0 for t in self._TOKENS]

    def embed_query(self, text: str) -> list[float]:
        return self._vec(text)

    def embed(self, texts: list[str]) -> list[list[float]]:
        return [self._vec(t) for t in texts]


# ------------------------------------------------------------------
# 11. Retriever returns standalone chunk
# ------------------------------------------------------------------


def test_retriever_returns_standalone_chunk():
    store = ProductStore.from_json(DATA)
    model = _KeywordModel()

    standalone = StandaloneDocument(
        source_file="docs/mdf_method.txt",
        kind="txt",
        title="Метод шлифовки МДФ",
        text="MDFTOKEN — уникальная метка для ретривал-теста.",
    )
    vs, _ = build_index(store, model, standalone_documents=[standalone])
    ret = Retriever(vector_store=vs, embedding_model=model)

    res = ret.search("MDFTOKEN question", top_k=5)
    standalone_hits = [
        rc for rc in res if rc.chunk.doc_type == "standalone"
    ]
    assert standalone_hits, "standalone chunk не найден в результатах"
    assert standalone_hits[0].chunk.source["file"] == "docs/mdf_method.txt"


# ------------------------------------------------------------------
# 12. ContextBuilder renders Title/Source for standalone chunks
# ------------------------------------------------------------------


def test_context_builder_renders_standalone_block():
    store = ProductStore.from_json(DATA)
    model = _KeywordModel()

    standalone = StandaloneDocument(
        source_file="docs/metal_method.txt",
        kind="txt",
        title="Метод: грунт под металл",
        text="UNIQUESTANDALONE999 — грунт под металл. Смешивание: 1:1.",
        section="металл",
    )
    vs, _ = build_index(store, model, standalone_documents=[standalone])
    ret = Retriever(vector_store=vs, embedding_model=model)
    cb = ContextBuilder(retriever=ret, product_store=store)

    ctx = cb.build(
        "UNIQUESTANDALONE999 how to prime metal",
        top_k=3,
        auto_detect_article=False,
    )
    assert ctx.has_context
    standalone_sources = [s for s in ctx.sources if s.doc_type == "standalone"]
    assert standalone_sources, "standalone source не представлен"
    ss = standalone_sources[0]
    assert ss.file == "docs/metal_method.txt"
    # rendered block mentions Title + Source (не Product)
    assert "Title: Метод: грунт под металл" in ctx.context
    assert "Source: docs/metal_method.txt" in ctx.context


# ------------------------------------------------------------------
# 13. create_rag_pipeline accepts standalone_root / standalone_documents
# ------------------------------------------------------------------


def test_create_rag_pipeline_standalone_documents(tmp_path: Path):
    (tmp_path / "notes.txt").write_text(
        "UNIQUE TOKEN — шлифование МДФ.\n",
        encoding="utf-8",
    )
    gen = create_rag_pipeline(
        products_path=DATA,
        llm=FakeLLM(answer="OK"),
        use_ollama=False,
        standalone_root=tmp_path,
    )
    assert isinstance(gen, AnswerGenerator)
    result = gen.answer(
        "UNIQUE TOKEN МДФ", top_k=3, auto_detect_article=False,
    )
    assert result.has_answer


def test_create_rag_pipeline_standalone_documents_kwarg():
    standalone = StandaloneDocument(
        source_file="docs/x.txt", kind="txt", text="Метод: шлифовка до P220",
    )
    gen = create_rag_pipeline(
        products_path=DATA,
        llm=FakeLLM(answer="OK"),
        use_ollama=False,
        standalone_documents=[standalone],
    )
    assert isinstance(gen, AnswerGenerator)


def test_load_standalone_documents_returns_docs(tmp_path: Path):
    (tmp_path / "a.txt").write_text("Применение: дерево\n", encoding="utf-8")
    (tmp_path / "junk.bin").write_bytes(b"x")
    docs = load_standalone_documents(tmp_path)
    # junk.bin → unsupported → goes to errors, NOT documents
    assert len(docs) == 1
    assert docs[0].source_file
    assert "дерево" in docs[0].text


# ------------------------------------------------------------------
# 14. End-to-end: LLM + ContextBuilder + standalone chunk
# ------------------------------------------------------------------


def test_e2e_fake_llm_with_standalone_context():
    store = ProductStore.from_json(DATA)
    model = _KeywordModel()

    standalone = StandaloneDocument(
        source_file="docs/sand_mdf.txt",
        kind="txt",
        title="Шлифование МДФ",
        text="UNIQUE999 — шлифовка МДФ до P180 обязательна перед грунтом.",
    )
    vs, _ = build_index(store, model, standalone_documents=[standalone])
    ret = Retriever(vector_store=vs, embedding_model=model)
    cb = ContextBuilder(retriever=ret, product_store=store)
    llm = FakeLLM(answer="Сначала шлифование МДФ, затем грунт.")
    gen = AnswerGenerator(context_builder=cb, llm=llm)

    result = gen.answer(
        "UNIQUE999 МДФ prepare",
        top_k=4,
        auto_detect_article=False,
    )
    assert result.has_answer
    # Prompt содержит standalone-блок (Title + Source).
    assert "Title: Шлифование МДФ" in llm.last_prompt
    assert "docs/sand_mdf.txt" in llm.last_prompt
