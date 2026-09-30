"""Task 10 — первичные источники совместимости (STAINWOOD).

Проверяет acceptance-критерии §7:
* обнаружение таблицы/схемы совместимости в первичном источнике;
* mapping артикулов из документа → Product/ProductVariant;
* импорт CONFIRMED в products.json;
* импорт FORBIDDEN в products.json;
* provenance у каждой связи;
* UNRESOLVED (неоднозначный mapping) НЕ попадает в compatibility;
* отсутствие документированной связи = UNKNOWN (а не FORBIDDEN);
* mixing-рецептура НЕ превращается в coating compatibility;
* runtime retrieval реально возвращает найденные compatibility relations.

Данные берутся из реальной ``data/knowledge/products.json`` (40 продуктов).
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from paint_rag.knowledge.compatibility_resolver import (
    ProductRef,
    make_resolver_from_products,
)
from paint_rag.knowledge.product_store import ProductStore
from paint_rag.models.product_relation import ProductRelation, RelationSource

ROOT = Path(__file__).resolve().parent.parent
KNOWLEDGE = ROOT / "data" / "knowledge"
PRODUCTS = KNOWLEDGE / "products.json"
LEGACY = KNOWLEDGE / "compatibility.json"
MATRIX = ROOT / "data/STAINWOOD/Схемы/Системы нанесения.xlsx"


@pytest.fixture(scope="module")
def store() -> ProductStore:
    return ProductStore.from_json(PRODUCTS)


@pytest.fixture(scope="module")
def product_dicts() -> list[dict]:
    return json.loads(PRODUCTS.read_text(encoding="utf-8"))


def _all_relations(store: ProductStore) -> list[ProductRelation]:
    out: list[ProductRelation] = []
    for p in store.products:
        out.extend(p.compatibility or [])
        for v in p.variants or []:
            out.extend(v.compatibility or [])
    return out


# ----------------------------------------------------------------------
# 1. Обнаружение первичного источника (таблица/схема совместимости)
# ----------------------------------------------------------------------


def test_primary_matrix_sheet_exists():
    wb = openpyxl.load_workbook(MATRIX, data_only=True)
    assert "Таблица1" in wb.sheetnames
    ws = wb["Таблица1"]
    header = (ws.cell(2, 1).value or "", ws.cell(2, 2).value or "")
    assert "Система" in header[0] or "систем" in header[0].lower()
    assert "Материалы" in header[1] or "соста" in header[1].lower()


def test_matrix_contains_documented_d2ur_system():
    """Первичный источник явным образом описывает Д-Дур систему
    (грунт + краска) — это и есть CONFIRMED-связка."""
    ws = openpyxl.load_workbook(MATRIX, data_only=True)["Таблица1"]
    found = None
    for r in range(3, ws.max_row + 1):
        name = ws.cell(r, 1).value
        materials = ws.cell(r, 2).value or ""
        if name and "Д-Дур" in str(name) and "краска" in str(materials):
            found = (name, materials)
            break
    assert found, "Д-Дур пигментированная система должна быть в Таблица1"
    name, materials = found
    assert "Д-Дур грунт" in materials
    assert "Д-Дур краска" in materials


# ----------------------------------------------------------------------
# 2. Mapping артикулов из документа → Product/ProductVariant
# ----------------------------------------------------------------------


def test_article_mapping_d2ur_topcoat(store: ProductStore):
    """Артикул 2675-755251 (Эмаль Д-Дур база 01) маппится на продукт."""
    p = store.get_by_article("2675-755251")
    assert p is not None
    assert "Эмаль" in p.name or "Д-ДУР" in p.name.upper()


def test_article_mapping_d2ur_primer(store: ProductStore):
    """Артикул 1149 (Д-Дур грунт) маппится на продукт (через артикул/алиас)."""
    primer = store.get_by_article("D-DUR 1149 (265-750001)")
    assert primer is not None
    # тот же продукт достижим по короткому алиасу «1149»
    assert primer in store.find("1149")
    assert "Д-ДУР" in primer.name.upper() or "Грунт" in primer.name


# ----------------------------------------------------------------------
# 3. Импорт CONFIRMED (артикул-уровень) в products.json
# ----------------------------------------------------------------------


def test_confirmed_d2ur_relation_imported(store: ProductStore):
    rel = next(
        (
            r
            for r in _all_relations(store)
            if r.base.upper() in {"1149", "265-750001"}
            and r.top.strip().upper() == "2675-755251"
            and r.status == "CONFIRMED"
        ),
        None,
    )
    assert rel is not None, "CONFIRMED Д-Дур грунт → Эмаль Д-Дур база 01 обязан быть"
    assert rel.allowed is True
    assert rel.source.has_provenance


def test_confirmed_relation_count(store: ProductStore):
    """В products.json есть несколько CONFIRMED-связей (не только Д-Дур)."""
    n_confirmed = sum(
        1 for r in _all_relations(store) if r.status == "CONFIRMED"
    )
    assert n_confirmed >= 10


# ----------------------------------------------------------------------
# 4. Импорт FORBIDDEN
# ----------------------------------------------------------------------


def test_forbidden_relation_imported(store: ProductStore):
    forb = [r for r in _all_relations(store) if r.status == "FORBIDDEN"]
    assert forb, "в products.json обязан быть хотя бы один FORBIDDEN"
    assert all(r.allowed is False for r in forb)
    assert all(r.source.has_provenance for r in forb)


def test_forbidden_bpd_sample(store: ProductStore):
    """Явный запрет «BPD-продукты не применять друг с другом»."""
    forb = [r for r in _all_relations(store) if r.status == "FORBIDDEN"]
    assert any("BPD" in (r.base + r.top).upper() for r in forb)


# ----------------------------------------------------------------------
# 5. Provenance обязательна
# ----------------------------------------------------------------------


def test_every_relation_has_provenance(store: ProductStore):
    rels = _all_relations(store)
    assert rels
    for r in rels:
        assert r.source.has_provenance, f"нет провенанса: {r}"
        assert (r.source.file or r.source.sheet or r.source.row is not None), (
            f"нет файла/листа/стrokes у связи {r.base}->{r.top}"
        )


def test_provenance_source_fields_roundtrip():
    src = RelationSource(file="a.pdf", page=1, note="цитата")
    out = src.as_dict()
    assert out["file"] == "a.pdf"
    assert out["page"] == 1
    assert src.has_provenance


# ----------------------------------------------------------------------
# 6. UNRESOLVED неоднозначный mapping НЕ импортирован
# ----------------------------------------------------------------------


def test_unresolved_mapping_not_imported(store: ProductStore):
    """«УС грунт» — второй грунт в цитате, но отдельного продукта нет:
    он НЕ становится relation'ом, а хранится как alternative."""
    all_names = {p.name for p in store.products}
    all_aliases = {a for p in store.products for a in p.aliases}
    has_us_grunt_product = any("уc грунт" in (n or "").lower() for n in all_names)
    has_us_grunt_alias = any(
        "уc грунт" in (a or "").lower() for a in all_aliases
    )
    assert not (has_us_grunt_product or has_us_grunt_alias), (
        "не должно быть продукта «УС грунт» как отдельной сущности"
    )
    # Но он сохранён как alternative у Эмали Д-Дур база 01.
    d2ur = store.get_by_article("2675-755251")
    assert d2ur is not None
    found_alt = any(
        "ус грунт" in (alt.ref or "").lower()
        for r in (d2ur.compatibility or [])
        for alt in (r.alternatives or [])
    )
    assert found_alt, "УС грунт должен быть сохранён как альтернатива праймера"


def test_unresolved_pu_paint_not_imported(store: ProductStore):
    """«ПУ-краска / ПУ-праймер / Краска Профф 355» — без артикула в
    products.json → НЕ создаются отдельные продукты ради связи."""
    names = {n.upper() for p in store.products for n in [p.name, *p.aliases]}
    assert not any("ПУ-КРАСКА" in n for n in names)
    assert not any("ПРОФФ" in n for n in names)


# ----------------------------------------------------------------------
# 7. Отсутствие связи = UNKNOWN (не FORBIDDEN)
# ----------------------------------------------------------------------


def test_absence_is_unknown_not_forbidden(store: ProductStore):
    """RUPA грунт PD155 → Лак PV210: в первичных источниках нет пары
    → resolver отвечает UNKNOWN, а не FORBIDDEN."""
    resolver = make_resolver_from_products(store.products)
    result = resolver.resolve(
        base=ProductRef(article="PD155"),
        top=ProductRef(article="PV210-XX"),
    )
    assert result.status == "UNKNOWN"
    assert result.status != "FORBIDDEN"


def test_no_unknown_relation_written(store: ProductStore):
    """UNKNOWN не пишется как явная запись: у PD155 нет связи с PV210
    в compatibility-списке."""
    p = next(x for x in store.products if "PD155" in x.name)
    for r in p.compatibility:
        assert not ("PD155" == r.base.upper() and "PV210" in r.top.upper())


# ----------------------------------------------------------------------
# 8. Mixing-рецептура ≠ coating compatibility
# ----------------------------------------------------------------------


def test_mixing_not_compatibility_pd155(store: ProductStore):
    p = next(x for x in store.products if "PD155" in x.name)
    # У варианта есть mixing-рецептура.
    v = p.variants[0]
    assert v.mixing is not None
    assert v.mixing.hardener is not None
    # Но mixing-компонент (HD820/HD810) НЕ стал coating-compatibility relation'ом.
    for r in _all_relations(store):
        if p.name in ("Грунт PD155 изолятор для МДФ",):
            assert "HD820" not in (r.base + r.top)
            assert "HD810" not in (r.base + r.top)


def test_alternatives_are_not_compatibility(store: ProductStore):
    """Альтернативные отвердители хранятся в `alternatives`, а не в
    `compatibility` (запрет §12 ТЗ)."""
    pd125 = next(x for x in store.products if "Грунт PD125" in x.name)
    assert any(a.ref == "HD820" for a in (pd125.alternatives or []))
    for r in pd125.compatibility:
        assert "HD820" not in (r.base + r.top)


# ----------------------------------------------------------------------
# 9. Legacy compatibility.json НЕ стал CONFIRMED (не авто-перенос)
# ----------------------------------------------------------------------


def test_legacy_chemistry_rules_not_promoted(product_dicts):
    """Legacy 25-правил NC/AC/PU/PE/UV/WB остаются legacy: ни один
    relation в products.json не имеет base/top уровня химии."""
    legacy = json.loads(LEGACY.read_text(encoding="utf-8"))
    chem_codes = {r["base"].upper() for r in legacy} | {
        r["top"].upper() for r in legacy
    }
    chem_codes = {c for c in chem_codes if c in {"NC", "AC", "PU", "PE", "UV", "WB"}}
    assert chem_codes, "legacy matrix codes expected"
    for p in product_dicts:
        for r in p.get("compatibility") or []:
            assert r["base"].upper() not in chem_codes, (
                f"legacy chemistry base {r['base']} не должен быть в products.json"
            )
            assert r["top"].upper() not in chem_codes, (
                f"legacy chemistry top {r['top']} не должен быть в products.json"
            )


# ----------------------------------------------------------------------
# 10. Runtime retrieval возвращает найденные compatibility relations
# ----------------------------------------------------------------------


def test_runtime_context_renders_compatibility_block(store: ProductStore):
    """documents._render_compatibility_block показывает CONFIRMED-связки
    с источником для Д-Дур грунта."""
    from paint_rag.rag.documents import _render_compatibility_block

    d2ur_primer = next(
        p for p in store.products if p.name == "Грунт Д-ДУР Плюс Белый"
    )
    rels = d2ur_primer.effective_compatibility(variant=None)
    assert rels
    block = "\n".join(_render_compatibility_block(rels))
    assert "Совместимость" in block
    assert "[CONFIRMED (совместимо)]" in block
    assert "base: 1149" in block
    assert "top: 2675-755251" in block
    assert "Источник" in block


def test_runtime_context_empty_unknown_when_no_relation(store: ProductStore):
    """Для продукта без совместимости блок рендерит UNKNOWN-фразу."""
    from paint_rag.rag.documents import _render_compatibility_block

    d2ur00 = next(p for p in store.products if p.name == "Эмаль Д-ДУР-00")
    block = "\n".join(_render_compatibility_block(d2ur00.compatibility))
    assert "UNKNOWN" in block or "не найдено" in block


def test_runtime_retrieval_confirmed_relation_present(store: ProductStore):
    """End-to-end: products → ProductStore → resolver видит CONFIRMED-пары
    (это то, что рендерится в CONTEXT для LLM и цитируется)."""
    resolver = make_resolver_from_products(store.products)
    result = resolver.resolve(
        base=ProductRef(article="1149"),
        top=ProductRef(article="2675-755251"),
    )
    assert result.status == "CONFIRMED"
    assert result.findings and result.findings[0].source.get("file")
